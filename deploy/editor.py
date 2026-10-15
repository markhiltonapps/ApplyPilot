"""Edit a tailored resume or cover letter in the browser, and download it as Word.

Saving rewrites the .txt, rebuilds the PDF next to it, and marks the resume as
checked so its review warning goes away.
"""

from __future__ import annotations

import html
import io
import json
from pathlib import Path

EDITABLE_DIRS = ("tailored_resumes", "cover_letters")

EDITOR_CSS = """
.editor textarea{width:100%;min-height:65vh;box-sizing:border-box;padding:14px;border-radius:12px;
  border:1px solid var(--line);background:var(--field);color:var(--text);
  font:14px/1.55 ui-monospace,SFMono-Regular,Menlo,Consolas,monospace;resize:vertical}
.editor .row{margin-top:12px}
"""


def resolve(data_dir: Path, rel: str) -> Path | None:
    """The .txt document a request names, only if it's inside an editable folder."""
    rel = (rel or "").strip().lstrip("/")
    if not rel.endswith(".txt") or rel.split("/", 1)[0] not in EDITABLE_DIRS:
        return None
    path = (data_dir / rel).resolve()
    root = (data_dir / rel.split("/", 1)[0]).resolve()
    return path if path.is_file() and path.is_relative_to(root) else None


def report_path(doc: Path) -> Path:
    return doc.with_name(doc.stem + "_REPORT.json")


def review_notes(doc: Path) -> str:
    """The judge's notes for a resume still awaiting a check, else ''."""
    try:
        data = json.loads(report_path(doc).read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return ""
    if data.get("status") != "approved_with_judge_warning":
        return ""
    issues = ((data.get("judge") or {}).get("issues") or "").strip()
    return "" if issues.lower() in ("", "none") else issues


def mark_checked(doc: Path) -> None:
    """Record that the person reviewed this resume, which clears its warning."""
    path = report_path(doc)
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return
    if data.get("status") == "approved_with_judge_warning":
        data["status"] = "approved"
        data["checked_by_user"] = True
        path.write_text(json.dumps(data, indent=2), encoding="utf-8")


def save(doc: Path, text: str) -> str | None:
    """Write the edited text and rebuild its PDF. Returns an error message if the PDF failed."""
    doc.write_text(text.replace("\r\n", "\n").strip() + "\n", encoding="utf-8")
    try:
        from applypilot.scoring.pdf import convert_to_pdf
        convert_to_pdf(doc)
    except Exception as e:  # noqa: BLE001 - the text is saved; report the PDF problem
        return f"Your changes are saved, but the PDF couldn't be rebuilt ({e})."
    return None


def page_body(doc: Path, rel: str, job_title: str, kind: str, message: str = "") -> str:
    notes = review_notes(doc) if kind == "resume" else ""
    notes_html = (f'<details class="review" open><summary>⚠ What to check</summary>'
                  f'<p>{html.escape(notes).replace(chr(10), "<br>")}</p></details>' if notes else "")
    msg = f'<div class="ok">{html.escape(message)}</div>' if message else ""
    checked_btn = ('<button class="btn ghost" name="action" value="checked" formnovalidate>'
                   'Looks fine, mark as checked</button>' if notes else "")
    q = html.escape(rel, quote=True)
    return f"""<p><a href="/ready">&larr; Ready to apply</a></p>
<h1>Edit {kind}</h1>
<p class="sub">{html.escape(job_title)}</p>
{msg}{notes_html}
<form class="editor" method="post" action="/edit">
<input type="hidden" name="f" value="{q}">
<textarea name="content" spellcheck="true">{html.escape(doc.read_text(encoding="utf-8"))}</textarea>
<p class="muted small">Keep section headings (SUMMARY, EXPERIENCE...) in capitals and bullets starting with
"- " so the PDF lays out correctly.</p>
<div class="row"><button class="btn primary" name="action" value="save">Save and rebuild PDF</button>
{checked_btn}
<a class="btn ghost" href="/docx?f={q}">Download as Word</a></div>
</form>"""


def to_docx(text: str) -> bytes:
    """A simple Word version of a resume or letter: name, capitalised headings, bullets."""
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    first = True
    for raw in text.splitlines():
        line = raw.rstrip()
        if not line.strip():
            continue
        stripped = line.strip()
        if first:
            run = doc.add_paragraph().add_run(stripped)
            run.bold, run.font.size = True, Pt(18)
            first = False
        elif stripped.isupper() and len(stripped) <= 40 and not stripped.startswith("-"):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            run = p.add_run(stripped)
            run.bold, run.font.size = True, Pt(11.5)
        elif stripped.startswith(("- ", "• ")):
            doc.add_paragraph(stripped[2:].strip(), style="List Bullet")
        else:
            doc.add_paragraph(stripped)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
